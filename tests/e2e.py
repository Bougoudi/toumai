"""Test de bout en bout de ProfitPilot (navigateur mobile + API + base Postgres)."""
import json, os, subprocess, time, urllib.request, urllib.error, urllib.parse, io, csv
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook

BASE = "http://localhost:8765"
RESULTS = []

def check(name, cond, info=""):
    RESULTS.append((bool(cond), name, info))
    print(("PASS " if cond else "FAIL ") + name + (f"  [{info}]" if info and not cond else ""))

def section(name, fn):
    print(f"\n== {name}")
    try:
        fn()
    except Exception as e:
        check(f"{name}: exécution sans exception", False, f"{type(e).__name__}: {str(e)[:300]}")

def sql(q):
    return subprocess.run(["su", "nobody", "-s", "/bin/sh", "-c",
        f"/usr/lib/postgresql/16/bin/psql -h /tmp/pp_e2e/pg -p 5499 -U pp postgres -tAc \"{q}\""],
        capture_output=True, text=True).stdout.strip()

def http(method, path, body=None, headers=None):
    data = json.dumps(body).encode() if body is not None else None
    h = {"Content-Type": "application/json", **(headers or {})}
    req = urllib.request.Request(BASE + urllib.parse.quote(path, safe="/?&=:"), data=data, method=method, headers=h)
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            raw = r.read()
            return r.status, (json.loads(raw) if raw[:1] in (b"{", b"[") else raw)
    except urllib.error.HTTPError as e:
        raw = e.read()
        try: return e.code, json.loads(raw)
        except Exception: return e.code, raw

CSV = ("Sipariş Numarası;Ürün Adı;Adet;Faturalanacak Tutar;Komisyon Tutarı;Maliyet;Sipariş Durumu\n"
       "1;Örnek Ürün;1;499;89,82;180;Teslim Edildi\n"
       "2;Pahalı Maliyet;2;200;36;500;Teslim Edildi\n"
       "3;Sıfır;1;0;0;10;Teslim Edildi\n")
os.makedirs("/tmp/pp_e2e", exist_ok=True)
open("/tmp/pp_e2e/r.csv", "w").write(CSV)

with sync_playwright() as p:
    b = p.chromium.launch(executable_path="/opt/pw-browsers/chromium-1194/chrome-linux/chrome")
    ctx = b.new_context(viewport={"width": 390, "height": 844}, accept_downloads=True)
    pg = ctx.new_page()
    errs = []
    pg.on("pageerror", lambda e: errs.append(str(e)))
    # les codes HTTP attendus (404, 422…) et les polices Google bloquées par le bac à sable ne sont pas des bugs
    pg.on("console", lambda m: errs.append("console: " + m.text) if m.type == "error" and "Failed to load resource" not in m.text else None)

    def landing():
        pg.goto(BASE + "/?src=instagram")
        check("landing : titre", "ProfitPilot" in pg.title())
        check("landing : pas de débordement horizontal (mobile)", pg.evaluate("document.documentElement.scrollWidth<=innerWidth+1"),
              pg.evaluate("document.documentElement.scrollWidth"))
        check("landing : calculateur net 103,66 ₺ (exemple par défaut)", pg.inner_text("#r-net") == "₺103,66", pg.inner_text("#r-net"))
        check("landing : stopaj 4,16 ₺", pg.inner_text("#r-stop") == "₺4,16", pg.inner_text("#r-stop"))
        pg.fill("#price", "100"); pg.dispatch_event("#price", "input")
        check("landing : produit à 100 ₺ en perte (verdict rouge)", "kaybettiriyor" in pg.inner_text("#verdict"))
        check("landing : 3 abonnements 249/499/999", all(x in pg.inner_text("#fiyatlar") for x in ["249 ₺", "499 ₺", "999 ₺"]))
        check("landing : section « Neden ProfitPilot? » (6 atouts)", pg.locator("#neden .step").count() == 6)
        pg.click('.plan-btn[data-plan="denetim"]'); pg.wait_for_timeout(700)
        check("landing : bouton audit pré-remplit l'offre", pg.input_value("#jPlan") == "denetim")
        pg.fill("#jName", "Zeynep"); pg.fill("#jStore", "Z Butik"); pg.fill("#jPhone", "0533 444 55 66"); pg.fill("#jSales", "600")
        pg.click("#jBtn"); pg.wait_for_function("document.getElementById('formMsg').textContent.includes('Talebin')")
        href = pg.get_attribute("#formMsg a", "href")
        check("landing : demande d'audit enregistrée + lien WhatsApp", "wa.me/905418624220" in href)
        check("landing : demande en base avec source instagram",
              sql("SELECT plan||'/'||source FROM waitlist WHERE phone='05334445566'") == "denetim/instagram")

    def gate_and_file():
        pg.goto(BASE + "/app.html?plan=pro")
        pg.wait_for_selector("#gate:not(.hidden)")
        check("app : écran d'inscription obligatoire", pg.is_visible("#gateNew"))
        pg.fill("#gName", "Ayşe"); pg.fill("#gStore", "Moda"); pg.fill("#gPhone", "12"); pg.click("#gBtn")
        pg.wait_for_function("document.getElementById('gMsg').textContent.includes('Geçerli')")
        check("app : numéro invalide refusé", True)
        pg.fill("#gPhone", "0532 111 22 33"); pg.click("#gBtn"); pg.wait_for_selector("#trialBar:not(.hidden)")
        check("app : essai démarré (14 jours)", "14 gün" in pg.inner_text("#trialBar"))
        check("app : plan choisi (pro) et source conservés", sql("SELECT plan||'/'||source FROM waitlist WHERE phone='05321112233'") == "pro/instagram",
              sql("SELECT plan||'/'||source FROM waitlist WHERE phone='05321112233'"))
        check("app : pas de débordement horizontal (mobile)", pg.evaluate("document.documentElement.scrollWidth<=innerWidth+1"))
        pg.wait_for_function("window.XLSX")
        pg.fill("#s-ads", "25"); pg.dispatch_event("#s-ads", "input")
        pg.set_input_files("#file", "/tmp/pp_e2e/r.csv"); pg.wait_for_selector("#tbody tr")
        check("fichier : import confirmé", pg.inner_text("#fileMsg").startswith("✅"))
        prods = pg.evaluate("Object.fromEntries(PRODUCTS.map(p=>[p.name,[+p.net_profit.toFixed(2),p.target_price&&+p.target_price.toFixed(2)]]))")
        check("calcul : net Örnek Ürün = 107,66 ₺ (vérifié à la main)", prods.get("Örnek Ürün", [0])[0] == 107.66, prods)
        check("calcul : prix cible 15 % = 436,51 ₺", prods.get("Örnek Ürün", [0, 0])[1] == 436.51, prods)
        check("calcul : ligne à 0 ₺ ignorée", "Sıfır" not in prods)
        check("qualité : fiabilité basse (coût > prix détecté)", "düşük" in pg.inner_text("#qualityBadge"))
        check("simulateur = tableau (net au prix actuel)", all(pg.evaluate("PRODUCTS.filter(p=>p.units).map(p=>Math.abs(netAt(p,p.avg_price)*p.units-p.net_profit)<0.01)")))
        with pg.expect_download() as d:
            pg.click("#xlsBtn")
        d.value.save_as("/tmp/pp_e2e/rapor.xlsx")
        wb = load_workbook("/tmp/pp_e2e/rapor.xlsx")
        check("Excel : 2 feuilles Özet + Ürünler", wb.sheetnames == ["Özet", "Ürünler"], wb.sheetnames)
        heads = [c.value for c in wb["Ürünler"][1]]
        check("Excel : colonne prix cible présente", any(h and h.startswith("Hedef fiyat") for h in heads), heads)

    def marketplaces():
        files = {
            # Amazon (rapport de commandes, tabulations, en anglais) : item-price = prix × quantité
            "amazon.txt": "order-id\torder-item-id\tsku\tproduct-name\tquantity-purchased\tcurrency\titem-price\titem-tax\tshipping-price\torder-status\n"
                          "111-1\ti1\tSKU1\tBlue Mug\t2\tTRY\t400\t0\t20\tShipped\n"
                          "111-2\ti2\tSKU1\tBlue Mug\t1\tTRY\t200\t0\t0\tCancelled\n"
                          "111-3\ti3\tSKU2\tRed Cup\t1\tTRY\t150\t0\t0\tShipped\n",
            "hepsiburada.csv": "Sipariş Numarası;Ürün Adı;Adet;Satış Fiyatı;Komisyon;Sipariş Durumu\n"
                               "H1;Kupa;2;250;90;Teslim Edildi\nH2;Kupa;1;250;45;İptal Edildi\nH3;Tabak;1;100;18;İade Edildi\n",
            "n11.csv": "Sipariş No;Ürün Adı;Miktar;Toplam Tutar;n11 Komisyonu;Durum\nN1;Çanta;3;900;162;Tamamlandı\n",
        }
        expected = {
            "amazon.txt": {"Blue Mug": [2, 400], "Red Cup": [1, 150]},
            "hepsiburada.csv": {"Kupa": [2, 500], "Tabak": [0, 0]},
            "n11.csv": {"Çanta": [3, 900]},
        }
        for fname, content in files.items():
            open("/tmp/pp_e2e/" + fname, "w", encoding="utf-8").write(content)
            pg.set_input_files("#file", "/tmp/pp_e2e/" + fname)
            pg.wait_for_function("document.getElementById('fileMsg').textContent.startsWith('✅')||document.getElementById('fileMsg').textContent.startsWith('⚠️')")
            got = pg.evaluate("Object.fromEntries(PRODUCTS.map(p=>[p.name,[p.units,p.revenue]]))")
            check(f"places de marché : {fname} lu correctement", got == expected[fname], got)
        check("places de marché : 14 choix dans la liste", pg.locator("#s-market option").count() == 14)

    def ai():
        pg.click("#aiBtn"); pg.wait_for_selector(".reco")
        check("IA : 3 recommandations affichées", pg.locator(".reco").count() == 3)
        with pg.expect_download() as d:
            pg.click("#xlsBtn")
        d.value.save_as("/tmp/pp_e2e/rapor2.xlsx")
        ozet = [r for r in load_workbook("/tmp/pp_e2e/rapor2.xlsx")["Özet"].iter_rows(values_only=True)]
        check("Excel : recommandations IA incluses", any(r[0] == "Yapay zekâ önerileri" for r in ozet))

    def support():
        pg.click("#helpBtn"); pg.fill("#helpQ", "Maliyet sütununu nasıl seçerim?"); pg.click("#helpSend")
        pg.wait_for_selector("#helpLog .acts button")
        check("assistant : répond et propose Çözüldü / Çözülmedi", "Sütun eşleştirme" in pg.inner_text("#helpLog"))
        pg.click("#helpLog .acts button:has-text('Çözülmedi')"); pg.wait_for_selector("#helpLog a.btn")
        check("assistant : « Çözülmedi » → ticket + WhatsApp", "talep #" in pg.inner_text("#helpLog"))
        pg.fill("#helpQ", "Aboneliği nasıl öderim?"); pg.click("#helpSend")
        pg.wait_for_function("document.querySelectorAll('#helpLog a.btn').length>=2")
        check("assistant : question paiement → transmise directement", sql("SELECT count(*) FROM tickets") == "2", sql("SELECT count(*) FROM tickets"))
        href = pg.locator("#helpLog a.btn").last.get_attribute("href")
        check("assistant : message WhatsApp contient nom et problème", "Ay%C5%9Fe" in href and "%C3%B6derim" in href)
        pg.click("#helpClose")

    def trendyol_flow():
        pg.click("#tyToggle"); pg.fill("#tySeller", "123456"); pg.fill("#tyKey", "KEYABCD"); pg.fill("#tySecret", "SECRETXYZ")
        pg.click("#tyBtn"); pg.wait_for_function("document.getElementById('fileMsg').textContent.startsWith('✅')", timeout=60000)
        prods = pg.evaluate("Object.fromEntries(PRODUCTS.map(p=>[p.name,[p.units,p.revenue,+p.commission.toFixed(2),p.returns]]))")
        check("Trendyol : annulée exclue, retour compté, commissions réelles",
              prods.get("Tişört") == [2, 600, 108, 0] and prods.get("Tayt") == [1, 400, 72, 1], prods)
        check("Trendyol : connexion enregistrée et chiffrée",
              sql("SELECT count(*) FROM connections") == "1" and "SECRETXYZ" not in sql("SELECT encode(creds,'escape') FROM connections"))
        check("Trendyol : clés effacées du formulaire", pg.input_value("#tyKey") == "" and pg.input_value("#tySecret") == "")
        pg.reload(); pg.wait_for_function("document.getElementById('fileMsg').textContent.startsWith('✅')", timeout=30000)
        check("Trendyol : données rechargées seules à l'ouverture", "son güncelleme" in pg.inner_text("#fileMsg"))
        code, _ = http("POST", "/api/cron/sync", headers={"X-Cron-Secret": "cronsecret"})
        time.sleep(3)
        check("nuit : synchro lancée (202) et état ok", code == 202 and sql("SELECT last_status FROM connections") == "ok")
        code, _ = http("POST", "/api/cron/sync", headers={"X-Cron-Secret": "faux"})
        check("nuit : mauvaise clé refusée (404)", code == 404)
        pg.click("#tyRefresh"); pg.wait_for_function("document.getElementById('fileMsg').textContent.includes('şimdi')", timeout=60000)
        check("Trendyol : « Şimdi güncelle » fonctionne", True)

    def stores():
        pg.fill("#s-ship", "45"); pg.dispatch_event("#s-ship", "input")
        pg.once("dialog", lambda d: d.accept("İkinci Mağaza")); pg.click("#storeAdd")
        check("magasins : nouveau magasin avec réglages par défaut", pg.input_value("#s-ship") == "60")
        check("magasins : pas de connexion Trendyol sur le nouveau", pg.eval_on_selector("#tyStatus", "e=>e.classList.contains('hidden')"))
        pg.select_option("#storeSel", "Mağazam"); pg.wait_for_timeout(1500)
        check("magasins : retour au 1er garde ses réglages et sa connexion",
              pg.input_value("#s-ship") == "45" and not pg.eval_on_selector("#tyStatus", "e=>e.classList.contains('hidden')"))

    def admin():
        adm = ctx.new_page()
        code, _ = http("GET", "/admin")
        check("admin : sans jeton → 404", code == 404)
        adm.goto(BASE + "/admin?token=x")
        txt = adm.inner_text("body")
        for part in ["Entonnoir", "Par source", "Connexions Trendyol", "Tickets d'aide"]:
            check(f"admin : section « {part} »", part in txt)
        check("admin : entonnoir compte l'essai démarré", "Essais gratuits démarrés" in txt)
        adm.locator("tr", has_text="Ayşe").first.locator("button").click(); adm.wait_for_load_state()
        check("admin : « +30 j » → payé", "payé jusqu'au" in adm.locator("tr", has_text="Ayşe").first.inner_text())
        code, raw = http("GET", "/admin/waitlist.csv?token=x")
        rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
        check("admin : export CSV avec colonnes essai/abonnement", "trial_at" in rows[0] and "paid_until" in rows[0], rows[0])
        adm.close()

    def expiry():
        sql("UPDATE waitlist SET trial_at = now() - interval '20 days', paid_until = NULL WHERE phone='05321112233'")
        pg.reload(); pg.wait_for_selector("#gateExpired:not(.hidden)")
        check("expiration : écran « Deneme süren doldu »", True)
        code, body = http("POST", "/api/analyze", {"token": pg.evaluate("localStorage.getItem('pp_trial')"),
            "products": [{"name": "x", "units": 1, "revenue": 1, "cost": 1, "commission": 0, "shipping": 0, "fees": 0,
                          "ads": 0, "returns_loss": 0, "vat": 0, "net_profit": 0, "margin_pct": 0}]})
        check("expiration : IA refusée (402)", code == 402, code)
        code, _ = http("POST", "/api/trendyol/refresh", {"token": pg.evaluate("localStorage.getItem('pp_trial')"), "store": "Mağazam"})
        check("expiration : synchro Trendyol refusée (402)", code == 402, code)

    def security():
        code, body = http("POST", "/api/trial", {"name": "Başkası", "store": "X", "phone": "05321112233"})
        check("sécurité : un numéro ne révèle pas le nom du client", body.get("name") == "Başkası", body.get("name"))
        code, _ = http("GET", "/api/trendyol/connection?token=faux&store=Mağazam")
        check("sécurité : connexion illisible sans jeton valide", code == 401, code)
        code, _ = http("POST", "/api/trendyol/sync", {"seller_id": "1", "api_key": "abcd", "api_secret": "abcd"})
        check("sécurité : synchro Trendyol sans essai → 401", code == 401, code)
        log = open("/tmp/pp_e2e/srv.log").read()
        check("sécurité : aucune clé Trendyol dans les journaux", "SECRETXYZ" not in log and "KEYABCD" not in log)

    for name, fn in [("Landing", landing), ("Essai + fichier", gate_and_file), ("Places de marché", marketplaces), ("IA", ai), ("Assistant", support),
                     ("Trendyol", trendyol_flow), ("Magasins", stores), ("Admin", admin),
                     ("Expiration", expiry), ("Sécurité", security)]:
        section(name, fn)
    check("aucune erreur JavaScript", not errs, errs[:5])
    b.close()

ok = sum(1 for r in RESULTS if r[0])
print(f"\nRÉSULTAT : {ok}/{len(RESULTS)} réussis")
for r in RESULTS:
    if not r[0]:
        print("  ✗", r[1], r[2])
